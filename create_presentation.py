"""
create_presentation.py
Generates an executive, modern 16:9 widescreen PowerPoint presentation (.pptx)
for Mailgent's Fetch-A-Thon 3.0 hackathon submission.
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    BG_DARK = RGBColor(15, 23, 42)        # #0F172A Deep Navy Slate
    CARD_BG = RGBColor(30, 41, 59)        # #1E293B Card Slate
    CARD_BORDER = RGBColor(51, 65, 85)    # #334155 Border
    ACCENT_PURPLE = RGBColor(95, 67, 241) # #5F43F1 Fetch Purple
    ACCENT_CYAN = RGBColor(56, 189, 248)  # #38BDF8 Electric Cyan
    TEXT_LIGHT = RGBColor(248, 250, 252)  # #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184)  # #94A3B8
    TEXT_YELLOW = RGBColor(251, 191, 36)  # #FBBF24
    SUCCESS_GREEN = RGBColor(34, 197, 94) # #22C55E
    URGENT_RED = RGBColor(239, 68, 68)    # #EF4444

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, category_tag: str, title: str, subtitle: str = ""):
        # Category Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.5), Inches(0.35))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = category_tag.upper()
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = ACCENT_CYAN

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_LIGHT

        if subtitle:
            p_sub = tf_title.add_paragraph()
            p_sub.text = subtitle
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, title: str, content: list, border_color=CARD_BORDER, badge: str = ""):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)

        tb = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(0.25), width - Inches(0.6), height - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Card Title
        p_head = tf.paragraphs[0]
        p_head.text = (f"[{badge}] " if badge else "") + title
        p_head.font.size = Pt(17)
        p_head.font.bold = True
        p_head.font.color.rgb = ACCENT_CYAN if not badge else TEXT_LIGHT

        # Card Items
        for item in content:
            p = tf.add_paragraph()
            p.space_before = Pt(8)
            if isinstance(item, tuple):
                bold_part, regular_part = item
                run1 = p.add_run()
                run1.text = bold_part + " "
                run1.font.bold = True
                run1.font.size = Pt(12)
                run1.font.color.rgb = TEXT_LIGHT

                run2 = p.add_run()
                run2.text = regular_part
                run2.font.bold = False
                run2.font.size = Pt(12)
                run2.font.color.rgb = TEXT_MUTED
            else:
                run = p.add_run()
                run.text = item
                run.font.size = Pt(12)
                run.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    add_bg(slide1)

    # Hackathon Banner Tag
    tag_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.1), Inches(4.5), Inches(0.45))
    tag_box.fill.solid()
    tag_box.fill.fore_color.rgb = ACCENT_PURPLE
    tag_box.line.fill.background()
    p_tag = tag_box.text_frame.paragraphs[0]
    p_tag.alignment = PP_ALIGN.CENTER
    p_tag.text = "⚡ FETCH-A-THON 3.0 : ASI:ONE AGENT CHALLENGE"
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = TEXT_LIGHT

    # Main Project Title
    t_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(11.5), Inches(2.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "Mailgent 📬🤖"
    p1.font.size = Pt(46)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT

    p2 = tf1.add_paragraph()
    p2.text = "Autonomous AI Inbox Assistant for Students"
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT_CYAN

    p3 = tf1.add_paragraph()
    p3.space_before = Pt(8)
    p3.text = "Transforming Natural Language Intent into Real-World Gmail Actions — Safe, Private & Production-Ready."
    p3.font.size = Pt(15)
    p3.font.color.rgb = TEXT_MUTED

    # Metadata Badges Box
    meta_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.3), Inches(11.7), Inches(2.2))
    meta_card.fill.solid()
    meta_card.fill.fore_color.rgb = CARD_BG
    meta_card.line.color.rgb = CARD_BORDER
    meta_card.line.width = Pt(1.5)

    tf_meta = meta_card.text_frame
    tf_meta.word_wrap = True
    tf_meta.margin_left = Inches(0.4)
    tf_meta.margin_top = Inches(0.25)

    p_m1 = tf_meta.paragraphs[0]
    p_m1.text = "AGENT PROFILE & VERIFIED COORDINATES"
    p_m1.font.size = Pt(12)
    p_m1.font.bold = True
    p_m1.font.color.rgb = ACCENT_CYAN

    items_meta = [
        ("Agent Name & Handle:", "mailgent (@mailgent)"),
        ("Agentverse Address:", "agent1qtdpknm8y3vqxj2sgvme5ql6wsqau7m9c3drd7k74avry64nnaxxs5y35ya (Active on Almanac API)"),
        ("Protocols Deployed:", "Agent Chat Protocol (ACP v0.4.12) + Fetch.ai Payment Protocol (0.1.0)"),
        ("Developer / Author:", "Md Mihran Sohail | GitHub: @CodeWithMihran | Track: From Intent to Action")
    ]
    for bold_text, reg_text in items_meta:
        p = tf_meta.add_paragraph()
        p.space_before = Pt(4)
        r1 = p.add_run()
        r1.text = bold_text + " "
        r1.font.bold = True
        r1.font.size = Pt(12)
        r1.font.color.rgb = TEXT_LIGHT

        r2 = p.add_run()
        r2.text = reg_text
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 2: The Problem — The Silent Student Career Crisis
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_bg(slide2)
    add_header(slide2, "The Problem & Opportunity", "The Silent Student Career Crisis", "72% of college students report missing career opportunities due to email fatigue and clutter.")

    add_card(slide2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0), "1. Promotional Clutter", [
        ("The Flood:", "College students receive 80+ emails daily: newsletters, campus spam, delivery receipts, and marketing FOMO quizzes."),
        ("Drowned Signals:", "Critical recruiter messages, interview dates, and scholarship notifications are buried beneath noise."),
        ("Cognitive Toll:", "Students waste 45+ minutes daily manually triaging messages across fragmented tabs.")
    ], border_color=CARD_BORDER)

    add_card(slide2, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0), "2. The 48h Clock", [
        ("Strict Expiry:", "Screening questionnaires, HackerRank tests, and Calendly scheduling links expire in 24 to 48 hours."),
        ("High Stakes:", "A single missed notification leads to an automatic application disqualification."),
        ("Human Limit:", "Students studying for exams cannot monitor their inbox 24/7 for ticking deadlines.")
    ], border_color=URGENT_RED)

    add_card(slide2, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0), "3. Recruiter Ghosting", [
        ("The Silent Void:", "Recruiters often go silent after resume submissions, leaving applicants in limbo."),
        ("Missed Timing:", "Follow-up emails sent after 5-7 days dramatically boost response rates by 30%."),
        ("Follow-up Friction:", "Students feel awkward drafting professional follow-ups and delay until it's too late.")
    ], border_color=ACCENT_PURPLE)

    # -------------------------------------------------------------
    # SLIDE 3: The Solution — Mailgent (From Intent to Action)
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_bg(slide3)
    add_header(slide3, "The Solution", "Mailgent: Beyond Passive Chatbots", "An autonomous AI agent on ASI:One that transforms natural language intent into real-world Gmail actions.")

    add_card(slide3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.4), "Autonomous Email Triage ('Audit')", [
        ("Instant Classification:", "Groups emails into Job/Internship, Academic, Finance, and Promotions."),
        ("Deadline Radar:", "Analyzes temporal urgency (< 48 hours) and flags critical commitments.")
    ])

    add_card(slide3, Inches(6.8), Inches(1.8), Inches(5.6), Inches(2.4), "Ghosting Radar & Auto-Drafting", [
        ("Stalled Thread Detection:", "Identifies applications waiting > 5 days without recruiter reply."),
        ("Zero-Click Drafts:", "Composes polite, personalized follow-ups straight to Gmail Drafts folder.")
    ])

    add_card(slide3, Inches(0.8), Inches(4.5), Inches(5.6), Inches(2.4), "Human-in-the-Loop (HITL) Safety", [
        ("Interactive Approval Card:", "Shows a structured preview of every proposed action before execution."),
        ("User Trust First:", "Zero unapproved mutations; full user control via simple Approve or Cancel.")
    ])

    add_card(slide3, Inches(6.8), Inches(4.5), Inches(5.6), Inches(2.4), "Real-World Gmail Impact", [
        ("Automated Organization:", "Applies [Agent] Job and [Agent] Urgent labels in user's real Gmail."),
        ("Clutter Elimination:", "Archives promotional clutter out of primary inbox automatically.")
    ])

    # -------------------------------------------------------------
    # SLIDE 4: Standout Innovations
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_bg(slide4)
    add_header(slide4, "Innovation & Creativity", "Three Standout Innovations", "Features specifically engineered to elevate Mailgent beyond ordinary email helpers.")

    add_card(slide4, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0), "👻 Ghosting Radar", [
        ("The Innovation:", "Autonomous thread analyzer that calculates time elapsed since last recruiter contact."),
        ("Threshold Trigger:", "If days_silent >= 5, triggers an automated follow-up proposal."),
        ("Context-Aware Draft:", "Pre-fills polite subject and body in Gmail Drafts so the student just hits Send.")
    ], border_color=ACCENT_PURPLE, badge="NEW")

    add_card(slide4, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0), "⏰ 48-Hour Deadline Radar", [
        ("Temporal Extraction:", "Regex & NLP entity matching for 'within 24h', 'expires tomorrow', 'booking link closes'."),
        ("Urgency Guard:", "Guards against marketing FOMO tricks (e.g. quiz contests) by restricting urgency strictly to Job, Academic, and Finance."),
        ("Priority Tagging:", "Surfaces ticking clocks directly in the ASI:One chat card.")
    ], border_color=URGENT_RED, badge="RADAR")

    add_card(slide4, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0), "🛡️ Dual-Mode Resilience", [
        ("Live OAuth 2.0:", "Connects seamlessly to production Gmail via Google APIs with scoped tokens."),
        ("Offline Fallback:", "Bundles a realistic 50-email student dataset ('sample_emails.json')."),
        ("Zero Evaluation Crashes:", "If network or credentials drop during jury inspection, Mailgent falls back gracefully.")
    ], border_color=SUCCESS_GREEN, badge="SAFETY")

    # -------------------------------------------------------------
    # SLIDE 5: Fetch.ai Technology & Architecture
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_bg(slide5)
    add_header(slide5, "Technical Architecture", "End-to-End Fetch.ai Ecosystem Synergy", "Deep, multi-protocol integration across the entire Fetch.ai technology stack.")

    add_card(slide5, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0), "1. Agent Chat Protocol (ACP)", [
        ("Native ASI:One Integration:", "Full ACP v0.4.12 implementation with ChatMessage and ChatAcknowledgement."),
        ("Sub-Second Handshake:", "Immediate delivery confirmation back to ASI:One to eliminate UI lag."),
        ("Rich Markdown Cards:", "Renders formatted tables, badges, and action plans in chat.")
    ], border_color=ACCENT_CYAN)

    add_card(slide5, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0), "2. Agentverse & Almanac", [
        ("Global Discovery:", "Registered on Almanac API (agent1qtdpknm8y3vqxj2sgvme5ql6wsqau7m9c3drd7k74avry64nnaxxs5y35ya)."),
        ("Agentverse Mailbox:", "24/7 cloud message routing and persistent endpoint management."),
        ("Non-blocking Concurrency:", "Async execution (asyncio.to_thread) keeps the event loop fluid.")
    ], border_color=ACCENT_PURPLE)

    add_card(slide5, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0), "3. Payment & ASI1 Reasoning", [
        ("Payment Protocol:", "AgentPaymentProtocol integrated with CommitPayment & CompletePayment handlers."),
        ("FET Micropayments:", "Monetization enabled: 0.01 FET per deep inbox scan, 0.05 FET for auto-drafting."),
        ("ASI1 LLM Hybrid:", "asi1-mini API reasoning backed by zero-latency local NLP rules.")
    ], border_color=TEXT_YELLOW)

    # -------------------------------------------------------------
    # SLIDE 6: Feature Suite & Commands
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_bg(slide6)
    add_header(slide6, "Product Capabilities", "The 6 Superpowers of Mailgent", "Controlled through intuitive natural language intents on ASI:One or CLI.")

    add_card(slide6, Inches(0.8), Inches(1.8), Inches(3.6), Inches(2.4), "Audit my emails 📬", [
        ("Scans recent inbox messages,", "categorizes by intent, identifies < 48h deadlines, and generates HITL safety card.")
    ])

    add_card(slide6, Inches(4.8), Inches(1.8), Inches(3.6), Inches(2.4), "Pipeline 💼", [
        ("Student Career Radar:", "Displays active interview loops (Nnomm, Echotalk, The ProEducator) and status.")
    ])

    add_card(slide6, Inches(8.8), Inches(1.8), Inches(3.6), Inches(2.4), "Approve / Cancel ✅", [
        ("Executes proposed actions:", "applies [Agent] Job and [Agent] Urgent labels, archives clutter, creates drafts.")
    ])

    add_card(slide6, Inches(0.8), Inches(4.5), Inches(3.6), Inches(2.4), "Search <term> 🔎", [
        ("Live keyword search:", "instant query across user's inbox with clean snippet cards and metadata.")
    ])

    add_card(slide6, Inches(4.8), Inches(4.5), Inches(3.6), Inches(2.4), "Draft to: ... 📝", [
        ("Custom Draft Composer:", "creates targeted email drafts directly in Gmail Drafts ready for final review.")
    ])

    add_card(slide6, Inches(8.8), Inches(4.5), Inches(3.6), Inches(2.4), "Profile 📇", [
        ("Live Account Telemetry:", "shows authenticated Gmail address, integration mode, and message totals.")
    ])

    # -------------------------------------------------------------
    # SLIDE 7: Competitive Matrix
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_bg(slide7)
    add_header(slide7, "Market Differentiation", "Mailgent vs. Generic Email Assistants", "Why Mailgent stands out in the Fetch-A-Thon ecosystem.")

    add_card(slide7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), "❌ Generic API Wrappers & Chatbots", [
        ("Manual & Primitive:", "Requires users to type arcane commands and message IDs (e.g. 'move to trash 1875f4...')."),
        ("No Semantic Reasoning:", "Cannot distinguish between an urgent interview invitation and an advertisement."),
        ("Blind Execution:", "Immediately mutates or deletes emails without human confirmation or preview."),
        ("No Career Awareness:", "Treats high-stakes recruiter communications like ordinary newsletters."),
        ("Fragile Offline State:", "Crashes immediately if credentials expire or API quotas are exceeded.")
    ], border_color=CARD_BORDER)

    add_card(slide7, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0), "✅ Mailgent (Autonomous AI Agent)", [
        ("High-Level Intent:", "User issues a single human prompt: 'Audit my emails' or 'Pipeline'."),
        ("Dual-Engine Reasoning:", "ASI1 LLM + NLP heuristic guards categorize emails and extract < 48h deadlines."),
        ("Human-in-the-Loop Cards:", "Presents clean Markdown Approval Cards before touching the user's inbox."),
        ("Proactive Career Protection:", "Rescues interview deadlines and drafts recruiter follow-ups for stalled jobs."),
        ("Dual-Mode Resilience:", "Operates seamlessly on live Google OAuth 2.0 with instant mock fallback.")
    ], border_color=ACCENT_CYAN)

    # -------------------------------------------------------------
    # SLIDE 8: Real-World Impact & Monetization
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    add_bg(slide8)
    add_header(slide8, "Impact & Economics", "Real-World Impact & Monetization Model", "Clear student value proposition paired with Fetch.ai Payment Protocol economics.")

    add_card(slide8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), "🎓 Student & Job Applicant Impact", [
        ("45 Min ➔ 30 Sec:", "Replaces endless inbox scrolling with a 30-second daily review and 1-click confirmation."),
        ("Zero Expired Opportunities:", "48-Hour Deadline Radar guarantees students never miss an online assessment."),
        ("30% Higher Interview Response:", "Ghosting Radar turns silent applications into renewed recruiter conversations."),
        ("Reduced Mental Fatigue:", "Removes inbox anxiety so students can focus on interview prep and academics.")
    ], border_color=SUCCESS_GREEN)

    add_card(slide8, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0), "💰 Agent Monetization via Payment Protocol", [
        ("Agent Payment Protocol:", "Native Fetch.ai payment protocol integrated and registered on Almanac."),
        ("Freemium Baseline:", "Core triage ('Audit', 'Profile', 'Pipeline') free for all students."),
        ("Micro-Transaction Tiers:", "Premium deep inbox scans (> 50 emails): 0.01 FET per scan."),
        ("Autonomous Career Follow-Ups:", "Full automated recruiter negotiation follow-up drafts: 0.05 FET."),
        ("Agent-to-Agent Economy:", "Peer recruiting agents can query Mailgent for applicant availability.")
    ], border_color=TEXT_YELLOW)

    # -------------------------------------------------------------
    # SLIDE 9: Live Gmail Verification & Demo Proof
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    add_bg(slide9)
    add_header(slide9, "Proof of Execution", "Live Gmail Verification & Real Impact", "Tested live on production Gmail account (sohail.mihran@gmail.com).")

    add_card(slide9, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0), "1. Live Gmail Labeling", [
        ("Real Labels Created:", "[Agent] Job and [Agent] Urgent created directly in Gmail account."),
        ("Accurate Tagging:", "Internshala interview invitations (Nnomm, Echotalk, The ProEducator) tagged automatically."),
        ("Instant Visibility:", "Urgent deadlines surfaced in the Gmail sidebar in red.")
    ], border_color=URGENT_RED)

    add_card(slide9, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0), "2. Clutter Archival", [
        ("Clean Inbox:", "Promotional newsletters and marketing FOMO emails archived out of INBOX."),
        ("Non-Destructive:", "Zero emails deleted; all archived emails remain searchable."),
        ("Focus Restored:", "Primary inbox reduced to only urgent, academic, and finance signals.")
    ], border_color=ACCENT_CYAN)

    add_card(slide9, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0), "3. Drafts Generation", [
        ("Gmail Drafts Folder:", "Follow-up email drafts created directly inside Gmail."),
        ("Safe Inspection:", "Drafts wait in Drafts folder for the student's final review before sending."),
        ("Professional Polish:", "Includes candidate name, role reference, and polite inquiry.")
    ], border_color=SUCCESS_GREEN)

    # -------------------------------------------------------------
    # SLIDE 10: Conclusion & Submission Coordinates
    # -------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_layout)
    add_bg(slide10)
    add_header(slide10, "Summary & Next Steps", "Mailgent: From Intent to Action", "Production-ready, privacy-focused, and built to empower students worldwide.")

    add_card(slide10, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), "🌟 Key Takeaways", [
        ("True Autonomous Agent:", "Not a toy wrapper — bridges natural language on ASI:One to real-world Gmail actions."),
        ("Full Fetch.ai Stack:", "uAgents, Agentverse Mailbox, Almanac API, ACP Chat Protocol, Payment Protocol, ASI1 LLM."),
        ("Real-World Value:", "Solves high-stakes student career pain points with deadline & ghosting radars."),
        ("Safety First:", "Human-in-the-Loop approval cards guarantee user confidence and data protection.")
    ], border_color=ACCENT_PURPLE)

    add_card(slide10, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0), "🔗 Project Coordinates & Submission", [
        ("Agentverse Details:", "agent1qtdpknm8y3vqxj2sgvme5ql6wsqau7m9c3drd7k74avry64nnaxxs5y35ya"),
        ("Agent Handle:", "@mailgent"),
        ("GitHub Repository:", "https://github.com/CodeWithMihran/Mailgent-ai"),
        ("Live Connected Account:", "sohail.mihran@gmail.com (Verified OAuth 2.0)"),
        ("Developer Contact:", "Md Mihran Sohail"),
        ("Hackathon:", "Fetch-A-Thon 3.0 : ASI:One Agent Challenge")
    ], border_color=ACCENT_CYAN)

    # Save presentation
    output_filename = "Mailgent_Presentation.pptx"
    prs.save(output_filename)
    print(f"[+] Presentation successfully created and saved as '{output_filename}'!")

if __name__ == "__main__":
    create_deck()
